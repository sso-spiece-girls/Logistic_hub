from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from core.excel_writer import CENTER_ALIGN

BBLUE = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
WHITE_F = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
NORM_F = Font(name="Calibri", size=10)
BOLD_F = Font(name="Calibri", size=10, bold=True)
THIN_B = Border(left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"))


def _hdr(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = BBLUE
        cell.font = WHITE_F
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_B


MESI = ["AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE 2026", "MAGGIO 2026",
        " GIUGNO 2026", "LUGLIO"]


class SoffasExcelWriter:
    def __init__(self, config):
        self.cfg = config

    def genera_excel(self, ddt_data_list, excel_path):
        wb = Workbook()
        self._crea_foglio1(wb, ddt_data_list)
        for mese in MESI:
            self._crea_mese(wb, mese)
        wb.save(excel_path)
        return excel_path

    def _crea_foglio1(self, wb, ddt_data_list):
        ws = wb.active
        ws.title = "Foglio1"
        headers = ["", "CODICE ARTICOLO", "NUMERO PACKING LIST",
                    "PARTITA", "LOTTO", "DATA ENTRATA", "DATA USCITA",
                    "", "TOT. BOBINE", ""]
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h)
        _hdr(ws, 1, 10)
        for col, w in zip("ABCDEFGHIJ", [6, 22, 18, 16, 10, 16, 16, 6, 12, 6]):
            ws.column_dimensions[col].width = w

        r = 2
        for ddt in ddt_data_list:
            extra = ddt.get("extra", {})
            partita = extra.get("partita", "")
            lotto = extra.get("lotto", "")
            packing = extra.get("packing_list", "")
            codice = ddt.get("ddt", "")
            data_entrata = ddt.get("data", "")
            bobine = extra.get("bobine", extra.get("totale_bobine", 6))
            for b_idx in range(int(bobine or 1)):
                ws.cell(row=r, column=1, value=r - 1).alignment = CENTER_ALIGN
                ws.cell(row=r, column=2, value=ddt.get("ddt", "")).alignment = CENTER_ALIGN
                ws.cell(row=r, column=3, value=packing).alignment = CENTER_ALIGN
                ws.cell(row=r, column=4, value=partita).alignment = CENTER_ALIGN
                ws.cell(row=r, column=5, value=lotto).alignment = CENTER_ALIGN
                ws.cell(row=r, column=6, value=data_entrata).alignment = CENTER_ALIGN
                r += 1
            ws.cell(row=r - 1, column=8, value="TOT. BOBINE")
            ws.cell(row=r - 1, column=9, value=bobine).alignment = CENTER_ALIGN

        ws.cell(row=r + 2, column=8, value="TOT AL 14/08").font = BOLD_F

    def _crea_mese(self, wb, nome_mese):
        ws = wb.create_sheet(nome_mese)
        ws.cell(row=1, column=1, value="Pallet Entrati")
        ws.cell(row=1, column=2, value="Pallet Usciti")
        ws.cell(row=1, column=4, value="Pallet In deposito")
        ws.cell(row=1, column=6, value="Prezzo")
        ws.cell(row=1, column=7, value="Deposito €")
        ws.cell(row=1, column=8, value="Ingresso a plt")
        ws.cell(row=1, column=9, value="Uscita a plt")
        _hdr(ws, 1, 9)

        ws.cell(row=3, column=4, value=0).alignment = CENTER_ALIGN
        ws.cell(row=3, column=6, value=6750).alignment = CENTER_ALIGN
        ws.cell(row=3, column=7, value=6.5).alignment = CENTER_ALIGN
        ws.cell(row=3, column=8, value=6).alignment = CENTER_ALIGN
        ws.cell(row=3, column=9, value=6).alignment = CENTER_ALIGN
        ws.cell(row=3, column=10, value="TARIFFE DA CONTRATTO")

        rate = 6
        ws.cell(row=4, column=1, value="=C33")
        ws.cell(row=4, column=2, value="=E33")
        ws.cell(row=4, column=4, value="=D3+A4-B4")
        ws.cell(row=4, column=6, value=rate).alignment = CENTER_ALIGN
        ws.cell(row=4, column=7, value=f"=F3*F4")
        ws.cell(row=4, column=8, value=6).alignment = CENTER_ALIGN
        ws.cell(row=4, column=9, value=6).alignment = CENTER_ALIGN

        ws.cell(row=7, column=2, value="Entrati plt n.")
        ws.cell(row=7, column=4, value="=C33")
        ws.cell(row=7, column=7, value=f"=D7*{rate}")
        ws.cell(row=8, column=2, value="Usciti plt n.")
        ws.cell(row=8, column=4, value="=E33")
        ws.cell(row=8, column=7, value=f"=D8*{rate}")
        ws.cell(row=9, column=2, value="EXTRA")
        ws.cell(row=9, column=7, value="=D9*22.5")
        ws.cell(row=10, column=7, value="=G4+G7+G8+G9")

        ws.cell(row=18, column=2, value="ENTRATI")
        ws.cell(row=18, column=4, value="USCITI")

        ws.cell(row=33, column=3, value="=SUM(C19:C32)")
        ws.cell(row=33, column=5, value="=SUM(E19:E32)")

        for col, w in zip("ABCDEFGHIJ", [10, 12, 12, 12, 12, 8, 12, 12, 12, 20]):
            ws.column_dimensions[col].width = w
